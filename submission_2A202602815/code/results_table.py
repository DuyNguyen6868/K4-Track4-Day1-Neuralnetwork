"""results_table.py — lưu kết quả từng lần chạy ra JSON, rồi điền vào experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx.

Tên cột của sheet "Experiments" (giữ nguyên, đúng thứ tự mẫu):
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
(các cột công thức ở cuối bảng mẫu tự tính, không ghi đè)
"""
from __future__ import annotations

import json
import math
from pathlib import Path

COLUMNS = ["exp_id", "group", "description", "loss", "optimizer", "lr", "weight_decay", "batch", "epochs",
           "hidden", "dropout", "clip_norm", "precision", "init", "seed", "step0_loss", "best_val_loss",
           "best_epoch", "final_train_loss", "final_val_loss", "val_acc", "val_macro_f1", "time_per_epoch_s",
           "peak_mem_MB", "diverged", "eval_acc", "eval_macro_f1", "figure_file", "notes"]
FORMULA_COLUMNS = {"step0_gap_vs_lnC", "gap_val_minus_train", "delta_val_f1_vs_base", "beyond_noise"}
MAX_ROWS = 60                      # công thức của mẫu chỉ phủ dòng 2..61
MAX_SEEDS = 5                      # sheet Seeds có 5 ô nhập (A2..A6)
GROUP_ORDER = ["baseline", "loss", "optimizer", "hparam", "dropout", "clipping", "amp", "init", "final", "other"]

# tên trong cfg -> chuỗi hiển thị mà bảng mẫu quy định
LOSS_NAMES = {"ce": "CE", "mse": "MSE"}
OPT_NAMES = {"sgd": "SGD", "sgd_momentum": "SGD+momentum", "adam": "Adam", "adamw": "AdamW"}


def _clean(obj):
    """NaN/inf -> None để file JSON hợp lệ; tuple -> list."""
    if isinstance(obj, dict):
        return {k: _clean(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [_clean(v) for v in obj]
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    return obj


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi cfg, history, summary (KHÔNG ghi best_state) ra <results_dir>/<exp_id>.json. Trả về đường dẫn."""
    out_dir = Path(results_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"{result['cfg']['exp_id']}.json"
    payload = _clean({k: result[k] for k in ("cfg", "history", "summary")})
    with open(path, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=1)
    return str(path)


def load_result(exp_id: str, results_dir: str = "../results") -> dict | None:
    """Đọc một file kết quả; trả về None nếu chưa có."""
    path = Path(results_dir) / f"{exp_id}.json"
    if not path.exists():
        return None
    with open(path, encoding="utf-8") as f:
        r = json.load(f)
    r["cfg"]["hidden"] = tuple(r["cfg"]["hidden"])
    return r


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict (sắp theo exp_id)."""
    ids = sorted(p.stem for p in Path(results_dir).glob("*.json"))
    return [load_result(i, results_dir) for i in ids]


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng: gộp cfg + summary (+ eval_acc, eval_macro_f1 nếu có)
    + figure_file. Chỉ truyền eval_scores (nội dung eval_result.json) cho baseline và cấu hình cuối cùng."""
    cfg, s = result["cfg"], result["summary"]
    all_notes = [n for n in (cfg.get("notes", ""), notes) if n]
    if cfg.get("scheduler"):
        all_notes.append(f"scheduler={cfg['scheduler']}")
    if s.get("peak_mem_MB") is None:
        all_notes.append(f"peak_mem_MB trống: chạy trên {s.get('device', 'cpu')}, không có GPU để đo")
    if s["diverged"]:
        all_notes.append("diverged: loss NaN/inf, dừng sớm")
    row = dict(
        exp_id=cfg["exp_id"], group=cfg["group"], description=cfg["description"],
        loss=LOSS_NAMES[cfg["loss"]], optimizer=OPT_NAMES[cfg["optimizer"]],
        lr=cfg["lr"], weight_decay=cfg["weight_decay"], batch=cfg["batch"], epochs=cfg["epochs"],
        hidden="-".join(str(h) for h in cfg["hidden"]), dropout=cfg["dropout"],
        clip_norm="none" if cfg["clip_norm"] is None else cfg["clip_norm"],
        precision=cfg["precision"], init=cfg["init"], seed=cfg["seed"],
        step0_loss=s["step0_loss"], best_val_loss=s["best_val_loss"], best_epoch=s["best_epoch"],
        final_train_loss=s["final_train_loss"], final_val_loss=s["final_val_loss"],
        val_acc=s["val_acc"], val_macro_f1=s["val_macro_f1"],
        time_per_epoch_s=s["time_per_epoch_s"], peak_mem_MB=s["peak_mem_MB"],
        diverged="Y" if s["diverged"] else "N",
        eval_acc=None if eval_scores is None else eval_scores["accuracy"],
        eval_macro_f1=None if eval_scores is None else eval_scores["macro_f1"],
        figure_file=f"figures/{cfg['exp_id']}.png", notes="; ".join(all_notes),
    )
    return row


def sort_rows(rows: list[dict]) -> list[dict]:
    """Sắp theo nhóm (thứ tự của bảng mẫu), trong nhóm giữ nguyên thứ tự đưa vào."""
    order = {g: i for i, g in enumerate(GROUP_ORDER)}
    return sorted(rows, key=lambda r: order.get(r["group"], len(order)))


def _cell_value(v):
    if v is None:
        return None
    if isinstance(v, float):
        return v if math.isfinite(v) else None
    return v


def write_xlsx(rows: list[dict], template_path: str, out_path: str,
               seed_ids: list[str] | None = None, summary_notes: dict[str, str] | None = None) -> None:
    """Điền các dòng vào sheet "Experiments" của mẫu (từ dòng 2), sheet "Seeds" (exp_id baseline) và cột
    nhận xét của sheet "Summary", rồi lưu thành out_path.

    openpyxl KHÔNG tính công thức: sau khi lưu phải mở file bằng Excel/LibreOffice rồi lưu lại để các ô
    xám (và sheet Seeds, Summary) có giá trị.
    """
    import openpyxl

    if len(rows) > MAX_ROWS:
        raise ValueError(f"bảng mẫu chỉ có công thức cho {MAX_ROWS} dòng, đang có {len(rows)}")
    ids = [r["exp_id"] for r in rows]
    assert len(set(ids)) == len(ids), "exp_id bị trùng"

    wb = openpyxl.load_workbook(template_path)          # không dùng data_only=True (sẽ mất công thức)
    ws = wb["Experiments"]
    col_of = {c.value: c.column for c in ws[1] if c.value}
    missing = [c for c in COLUMNS if c not in col_of]
    assert not missing, f"bảng mẫu thiếu cột {missing}"
    for i, row in enumerate(rows):
        r = i + 2
        for name in COLUMNS:                            # chỉ ghi cột dữ liệu, bỏ qua FORMULA_COLUMNS
            ws.cell(row=r, column=col_of[name]).value = _cell_value(row.get(name))

    if seed_ids is not None:
        if len(seed_ids) > MAX_SEEDS:
            raise ValueError(f"sheet Seeds chỉ nhận {MAX_SEEDS} seed")
        wss = wb["Seeds"]
        for i in range(MAX_SEEDS):
            wss.cell(row=i + 2, column=1).value = seed_ids[i] if i < len(seed_ids) else None

    if summary_notes:
        wsm = wb["Summary"]
        head = {c.value: c.column for c in wsm[1] if c.value}
        note_col = head["nhận xét ngắn (bạn viết)"]
        for r in range(2, wsm.max_row + 1):
            g = wsm.cell(row=r, column=1).value
            if g in summary_notes:
                wsm.cell(row=r, column=note_col).value = summary_notes[g]

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
